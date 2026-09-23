"""ตรวจโปรไฟล์เกียรติบัตรกับ "ไฟล์จริง" ที่วางไว้ใน apps/worker/tmp/

รันในคอนเทนเนอร์:
    docker compose exec worker python scripts/check_real_files.py

หาโฟลเดอร์ที่ตั้งชื่อตามรายการและรอบเอง เช่น tmp/Cert/HKIMO_/FINAL/ หรือ tmp/HKIMO_HEAT_2026/
แล้วอ่านทุกหน้าด้วยโปรไฟล์ของรายการ/รอบนั้น รายงานว่าสกัดข้อมูลได้ครบแค่ไหน
ถ้ามีไฟล์ Excel รายชื่อ (.xlsx) ที่ชื่อบอกรายการและรอบ จะเทียบเลขกับชื่อให้ด้วย

สคริปต์นี้ไม่แตะฐานข้อมูลและไม่อัปโหลดอะไร
ไฟล์จริงมีข้อมูลส่วนบุคคล จึง **ไม่ถูก commit เข้า git** (tmp/ อยู่ใน .gitignore)
ถ้าไม่มีไฟล์ สคริปต์จะบอกแล้วจบอย่างสงบ ไม่ถือว่าล้มเหลว
"""

import collections
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pymupdf  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.certificate_profiles import (  # noqa: E402
    CertificateProfile,
    Nationality,
    load_manifests,
    get_profile,
)
from app.normalize import normalize_name  # noqa: E402
from app.tasks.extract import page_text  # noqa: E402

TMP = Path("/app/tmp")
ROUNDS = ("HEAT", "FINAL")


def main() -> int:
    if not TMP.exists():
        print(f"ไม่พบ {TMP} — ข้ามการตรวจ (ปกติ ไฟล์จริงไม่ได้อยู่ใน git)")
        return 0

    targets = _discover()
    if not targets:
        print(f"ไม่พบโฟลเดอร์ที่ตั้งชื่อตามรายการ/รอบใน {TMP} — ข้ามการตรวจ")
        return 0

    ok = True
    for program, exam_round, folder in targets:
        profile = get_profile(program, exam_round)
        pages = _read(folder, profile)
        if not pages:
            continue
        ok = _report(profile, folder, pages) and ok
        roster = _matching_roster(program, exam_round)
        if roster:
            ok = _report_roster(pages, roster) and ok

    print()
    print("ผ่านทั้งหมด ✓" if ok else "มีข้อที่ไม่ผ่าน ✗")
    return 0 if ok else 1


def _discover() -> list[tuple[str, str, Path]]:
    """หาโฟลเดอร์ <รายการ>/<รอบ>/ หรือ <รายการ>_<รอบ>_<ปี>/ ที่มี PDF อยู่ข้างใน"""
    programs = {p for p, _ in load_manifests().supported_pairs()}
    found: list[tuple[str, str, Path]] = []
    for folder in sorted(p for p in TMP.rglob("*") if p.is_dir()):
        name = folder.name.upper()
        parent = folder.parent.name.upper().strip("_ ")
        match = re.fullmatch(r"([A-Z]+)_(HEAT|FINAL)(?:_\d{4})?", name)
        if name in ROUNDS and parent in programs:
            found.append((parent, name, folder))
        elif match and match.group(1) in programs:
            found.append((match.group(1), match.group(2), folder))
    return [t for t in found if any(t[2].rglob("*.pdf"))]


def _read(folder: Path, profile: CertificateProfile) -> list[dict]:
    """อ่านทุกหน้าในโฟลเดอร์ — รางวัลจากชื่อโฟลเดอร์ที่ใกล้ไฟล์ที่สุด (เครื่องมือตรวจเท่านั้น)"""
    pages: list[dict] = []
    for path in sorted(folder.rglob("*.pdf")):
        award = next(
            (a for part in reversed(path.relative_to(folder).parts[:-1])
             if (a := profile.catalog.resolve_folder(part))),
            None,
        )
        if award is None:
            print(f"  ! {path.relative_to(TMP)} — ไม่ได้อยู่ในโฟลเดอร์รางวัลที่ {profile.key} รู้จัก")
            continue
        with pymupdf.open(path) as doc:
            for i in range(doc.page_count):
                parsed = profile.parse(page_text(doc[i]))
                pages.append({"award": award, "parsed": parsed, "nationality": profile.nationality(parsed)})
    return pages


def _report(profile: CertificateProfile, folder: Path, pages: list[dict]) -> bool:
    total = len(pages)
    print("=" * 72)
    print(f"{profile.key}: {folder.relative_to(TMP)} ({total} หน้า)")
    print("=" * 72)

    parsed = [p["parsed"] for p in pages]
    checks = {
        "อ่านชื่อได้": sum(1 for p in parsed if p.name),
        "อ่านเลขผู้เข้าสอบได้": sum(1 for p in parsed if p.candidate_no),
        "อ่านระดับชั้นได้": sum(1 for p in parsed if p.level),
        "อ่านปีได้": sum(1 for p in parsed if p.year_on_page),
    }
    ok = True
    for label, count in checks.items():
        mark = "✓" if count == total else "✗"
        ok = ok and count == total
        print(f"  {mark} {label}: {count}/{total}")

    from_label = "โรงเรียน" if profile.round == "HEAT" else "ประเทศ"
    with_from = sum(1 for p in parsed if p.school_on_page or p.country_on_page)
    print(f"  · อ่านค่าหลัง from ({from_label}) ได้: {with_from}/{total}")
    nationality = collections.Counter(p["nationality"].value for p in pages)
    print(f"  · สัญชาติ: {dict(nationality)}")
    if nationality.get(Nationality.UNVERIFIED.value):
        print("    (หน้าที่ไม่มีหลักฐานสัญชาติจะถูกส่งให้แอดมินยืนยัน ไม่ถูกทิ้งเงียบ ๆ)")

    by_award = collections.Counter(p["award"].code for p in pages)
    print(f"  · รางวัลจากโฟลเดอร์: {dict(by_award)}")
    with_text = [p for p in pages if p["parsed"].award_text]
    mismatched = [
        p for p in with_text
        if (found := profile.catalog.resolve_text(p["parsed"].award_text)) is None or found.code != p["award"].code
    ]
    print(f"  {'✓' if not mismatched else '✗'} ข้อความรางวัลบนหน้าตรงกับโฟลเดอร์: "
          f"{len(with_text) - len(mismatched)}/{len(with_text)} (หน้าที่ไม่มีบรรทัดรางวัลเป็นปกติ)")
    ok = ok and not mismatched

    errors = [p for p in parsed if p.errors]
    print(f"  {'✓' if not errors else '✗'} รอบ/ปีบนหน้าขัดกันเอง: {len(errors)} หน้า")
    extras = collections.Counter(k for p in parsed for k in p.extra)
    if extras:
        print(f"  · ข้อมูลเฉพาะรายการที่อ่านได้: {dict(extras)}")

    numbers = [p.candidate_no for p in parsed if p.candidate_no]
    print(f"  · เลขไม่ซ้ำ {len(set(numbers))} เลข จาก {len(numbers)} หน้า "
          f"({len(numbers) - len(set(numbers))} หน้าใช้เลขซ้ำ = คนเดียวได้หลายใบ)")
    return ok


def _matching_roster(program: str, exam_round: str) -> Path | None:
    """ไฟล์ Excel ที่ชื่อบอกรายการและรอบ เช่น 'HKIMO Final 2026 - TH-1.xlsx'"""
    for path in sorted(TMP.glob("*.xlsx")):
        name = path.name.upper()
        if program in name and exam_round in name:
            return path
    return None


def _report_roster(pages: list[dict], xlsx_path: Path) -> bool:
    print(f"  เทียบกับไฟล์ Excel: {xlsx_path.name}")
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    rows = [r for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)][1:]
    roster = {str(r[0]).strip(): str(r[2]).strip() for r in rows if r and r[0] is not None}

    thai = [p["parsed"] for p in pages if p["nationality"] is Nationality.ACCEPT]
    found = [p for p in thai if p.candidate_no in roster]
    name_ok = [p for p in found if normalize_name(roster[p.candidate_no]) == normalize_name(p.name or "")]
    ok = len(found) == len(thai) and len(name_ok) == len(found)
    print(f"  {'✓' if len(found) == len(thai) else '✗'} หน้าที่หาเลขเจอใน Excel: {len(found)}/{len(thai)}")
    print(f"  {'✓' if len(name_ok) == len(found) else '✗'} เลขตรงและชื่อตรงด้วย: {len(name_ok)}/{len(found)}")
    unused = set(roster) - {p.candidate_no for p in thai}
    print(f"  {'✓' if not unused else '!'} แถว Excel ที่ไม่มีหน้าเกียรติบัตร: {len(unused)}")
    return ok


if __name__ == "__main__":
    raise SystemExit(main())
