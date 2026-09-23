"""รายชื่อผู้เข้าสอบ (roster) — ตรวจไฟล์ Excel เป็นร่าง แล้วสลับเป็นชุดที่ใช้อยู่เมื่อแอดมินสั่ง

รายชื่อคือแหล่งความจริงว่า "ใครควรได้เกียรติบัตร" ต้องมาก่อนอัป ZIP เสมอ
ไฟล์เดียวมีทั้งผู้เข้าสอบ online และ onsite โดยแต่ละแถวต้องบอกรูปแบบการสอบ

สองขั้นแยกกันโดยตั้งใจ:
  1. ตรวจ (ROSTER_VALIDATE) — อ่านไฟล์ ตรวจทุกแถว เก็บเป็นร่าง **ไม่แตะรายชื่อที่ใช้อยู่เลย**
     ไฟล์ผิดตรงไหนรายงานครบทุกแถว ไม่หยุดที่แถวแรก
  2. ใช้ (ROSTER_ACTIVATE) — แอดมินเห็นยอดรวมแล้วกดใช้ จึงสลับในทรานแซกชันเดียว
     รายการที่แอดมินเพิ่มเอง (MANUAL) รอดเสมอ ถ้าชนกับแถวใหม่ต้องให้แอดมินตัดสินก่อน
"""

from __future__ import annotations

import io
import json
import logging
import re
from dataclasses import dataclass, field
from typing import Any, Callable

from openpyxl import load_workbook

from ..audit import record_audit
from ..db import connection, new_id
from ..normalize import name_sort_key, normalize_name, normalize_school
from ..storage import download_bytes

log = logging.getLogger(__name__)

ProgressFn = Callable[[dict[str, Any]], None]

MODES = ("ONLINE", "ONSITE")

# หัวคอลัมน์ที่ยอมรับ — เทียบหลังตัดช่องว่าง จุด ขีดล่าง และทำตัวพิมพ์เล็กแล้ว
HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    # ของจริงใช้ CANDIDATE NO ซึ่งเป็นเลขเดียวกับ "Cert No" ที่พิมพ์บนหน้าเกียรติบัตร
    "cert_no": ("candidateno", "candidatenumber", "certno", "certificateno", "certificatenumber",
                "เลขที่นั่งสอบ", "เลขประจำตัวสอบ", "รหัสผู้เข้าสอบ", "เลขเกียรติบัตร",
                "เลขที่เกียรติบัตร", "เลขผู้เข้าสอบ"),
    "name_en": ("candidatename", "name", "fullname", "englishname", "nameenglish", "nameen",
                "nameeng", "ชื่อภาษาอังกฤษ", "ชื่ออังกฤษ", "ชื่อ-นามสกุลภาษาอังกฤษ"),
    "first_en": ("firstname", "givenname", "given"),
    "last_en": ("lastname", "surname", "familyname", "family"),
    "name_th": ("ชื่อ-นามสกุล", "ชื่อ-สกุล", "ชื่อนามสกุล", "ชื่อภาษาไทย", "ชื่อไทย",
                "ชื่อ-นามสกุลภาษาไทย", "thainame"),
    "first_th": ("ชื่อ", "ชื่อจริง"),
    "last_th": ("นามสกุล", "สกุล"),
    "level": ("grade", "ระดับชั้น", "ชั้น", "class"),
    "award": ("รางวัล", "ผลการแข่งขัน", "ผลรางวัล", "award", "result", "medal", "prize"),
    "school": ("โรงเรียน", "สถานศึกษา", "ชื่อโรงเรียน", "school", "schoolname", "institution"),
    # รูปแบบการสอบ — บังคับต้องมี เพราะไฟล์เดียวมีทั้ง online และ onsite
    "exam_mode": ("exammode", "mode", "testmode", "sittingmode", "online/onsite", "onsite/online",
                  "onlineonsite", "รูปแบบการสอบ", "รูปแบบสอบ", "ประเภทการสอบ", "ช่องทางการสอบ",
                  "สอบแบบ"),
}

# ชื่อคอลัมน์ที่แสดงในข้อความผิดพลาด — ใช้ชื่อแบบที่ไฟล์ของจริงเขียน
COLUMN_LABELS = {
    "cert_no": "CANDIDATE NO",
    "name": "CANDIDATE NAME",
    "exam_mode": "EXAM MODE",
}

MAX_HEADER_SCAN_ROWS = 10
# เก็บรายละเอียดข้อผิดพลาดไว้ไม่เกินนี้ — เกินกว่านี้เก็บแค่จำนวน (ไฟล์ผิดทั้งคอลัมน์มีเป็นพันแถว)
MAX_ERRORS = 500


@dataclass(frozen=True)
class RosterRow:
    row_number: int
    candidate_no: str
    name_en: str
    name_th: str
    exam_mode: str
    level: str
    school: str
    award: str


@dataclass
class RosterValidation:
    rows: list[RosterRow] = field(default_factory=list)
    errors: list[dict[str, Any]] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def error(self, message: str, row: int | None = None, column: str | None = None) -> None:
        self.errors.append({"row": row, "column": column, "message": message})

    def counts(self) -> dict[str, int]:
        return {
            "total": len(self.rows),
            "online": sum(1 for r in self.rows if r.exam_mode == "ONLINE"),
            "onsite": sum(1 for r in self.rows if r.exam_mode == "ONSITE"),
        }


# ---------------------------------------------------------------- อ่านและตรวจไฟล์


def validate_roster(xlsx_bytes: bytes) -> RosterValidation:
    """อ่านไฟล์และตรวจทุกแถว — คืนแถวที่ผ่านพร้อมข้อผิดพลาดทั้งหมด

    ไฟล์ใช้ได้ก็ต่อเมื่อไม่มีข้อผิดพลาดเลยแม้แต่แถวเดียว (ปฏิเสธทั้งร่าง ไม่ใช่ข้ามแถวที่ผิด)
    เพราะแถวที่หายไปเงียบ ๆ = เด็กคนนั้นไม่ได้เกียรติบัตรโดยไม่มีใครรู้
    """
    result = RosterValidation()
    try:
        workbook = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    except Exception:  # ไฟล์เสีย หรือไม่ใช่ .xlsx — openpyxl โยนได้หลายชนิด
        result.error("เปิดไฟล์ไม่ได้ ต้องเป็นไฟล์ Excel (.xlsx)")
        return result

    grid = [list(r) for r in workbook.active.iter_rows(values_only=True)]
    header_index, columns = _find_header(grid)
    if header_index is None:
        result.error(
            "หาหัวตารางไม่เจอ — ต้องมีคอลัมน์ชื่อผู้เข้าสอบอย่างน้อย 1 คอลัมน์ "
            "ในสิบแถวแรก (เช่น CANDIDATE NAME)"
        )
        return result

    for required in ("cert_no", "exam_mode"):
        if required not in columns:
            result.error(f"ไม่มีคอลัมน์ {COLUMN_LABELS[required]}", column=COLUMN_LABELS[required])
    if result.errors:
        return result

    rows_by_no: dict[str, list[int]] = {}
    for row_number, raw in enumerate(grid[header_index + 1:], start=header_index + 2):
        if all(cell is None or str(cell).strip() == "" for cell in raw):
            continue  # แถวว่างล้วน — ไม่ใช่ข้อผิดพลาด
        row = _build_row(raw, columns, row_number, result)
        if row:
            result.rows.append(row)
            rows_by_no.setdefault(row.candidate_no, []).append(row_number)

    for candidate_no, row_numbers in rows_by_no.items():
        if len(row_numbers) > 1:
            listed = ", ".join(str(n) for n in row_numbers)
            for row_number in row_numbers:
                result.error(
                    f"เลขผู้เข้าสอบ {candidate_no} ซ้ำกันในแถว {listed} "
                    "(เลขต้องไม่ซ้ำกันทั้งไฟล์ ทั้ง online และ onsite)",
                    row=row_number, column=COLUMN_LABELS["cert_no"],
                )

    if len(result.errors) > MAX_ERRORS:
        extra = len(result.errors) - MAX_ERRORS
        result.errors = result.errors[:MAX_ERRORS]
        result.error(f"และยังมีปัญหาอื่นอีก {extra} รายการ")
    return result


def _build_row(
    raw: list[Any], columns: dict[str, int], row_number: int, result: RosterValidation
) -> RosterRow | None:
    def value(field_name: str) -> str:
        col = columns.get(field_name)
        if col is None or col >= len(raw) or raw[col] is None:
            return ""
        return str(raw[col]).strip()

    name_en = value("name_en") or " ".join(x for x in (value("first_en"), value("last_en")) if x)
    name_th = value("name_th") or " ".join(x for x in (value("first_th"), value("last_th")) if x)
    raw_no = value("cert_no")
    raw_mode = value("exam_mode")

    problems_before = len(result.errors)
    candidate_no = _candidate_no(raw_no)
    if not raw_no:
        result.error("ไม่มีเลขผู้เข้าสอบ", row_number, COLUMN_LABELS["cert_no"])
    elif not candidate_no:
        result.error(f"เลขผู้เข้าสอบ '{raw_no}' ต้องเป็นตัวเลขล้วน", row_number, COLUMN_LABELS["cert_no"])

    if not normalize_name(name_en) and not normalize_name(name_th):
        result.error("ไม่มีชื่อผู้เข้าสอบที่ใช้ได้", row_number, COLUMN_LABELS["name"])

    mode = parse_mode(raw_mode)
    if not raw_mode:
        result.error("ไม่ได้ระบุรูปแบบการสอบ (ONLINE หรือ ONSITE)", row_number, COLUMN_LABELS["exam_mode"])
    elif mode is None:
        result.error(
            f"รูปแบบการสอบ '{raw_mode}' ไม่ถูกต้อง ต้องเป็น ONLINE หรือ ONSITE เท่านั้น",
            row_number, COLUMN_LABELS["exam_mode"],
        )

    if len(result.errors) > problems_before:
        return None
    return RosterRow(
        row_number=row_number,
        candidate_no=candidate_no,
        name_en=name_en,
        name_th=name_th,
        exam_mode=mode,
        level=value("level"),
        school=value("school"),
        award=value("award"),
    )


def parse_mode(raw: str) -> str | None:
    """ตัดช่องว่างหัวท้ายและไม่สนตัวพิมพ์ แต่รับแค่ ONLINE กับ ONSITE เท่านั้น

    'On-site' หรือ 'ออนไลน์' ไม่ผ่าน — ถ้ายอมเดาคำสะกดอื่น วันหนึ่งจะเดาผิด
    """
    value = (raw or "").strip().upper()
    return value if value in MODES else None


def _candidate_no(text: str) -> str:
    """เลขผู้เข้าสอบ — คืนค่าว่างถ้าไม่ใช่ตัวเลขล้วน

    ต้องตัดส่วนทศนิยมทิ้งก่อน เพราะ openpyxl อ่านเซลล์ตัวเลขมาเป็น float
    แล้ว str() ได้ '900101.0' — ถ้าเก็บแค่ตัวเลขตรง ๆ จะกลายเป็น '9001010'
    ซึ่งทำให้เลขผู้เข้าสอบเพี้ยนทั้งไฟล์โดยไม่มีอะไรฟ้อง
    """
    cleaned = re.sub(r"\.0+$", "", text.strip())
    return cleaned if cleaned.isdigit() else ""


def _find_header(grid: list[list[Any]]) -> tuple[int | None, dict[str, int]]:
    for index, raw in enumerate(grid[:MAX_HEADER_SCAN_ROWS]):
        columns: dict[str, int] = {}
        for col, cell in enumerate(raw):
            name = _match_header(cell)
            # คอลัมน์แรกที่เจอชนะ กันกรณีมีหัวซ้ำ
            if name and name not in columns:
                columns[name] = col
        if _has_name_column(columns):
            return index, columns
    return None, {}


def _match_header(cell: Any) -> str | None:
    if cell is None:
        return None
    key = "".join(str(cell).split()).lower().replace("_", "").replace(".", "")
    if not key:
        return None
    for name, aliases in HEADER_ALIASES.items():
        if key in aliases:
            return name
    return None


def _has_name_column(columns: dict[str, int]) -> bool:
    return bool(
        {"name_en", "name_th"} & columns.keys()
        or {"first_en", "last_en"} <= columns.keys()
        or {"first_th", "last_th"} <= columns.keys()
    )


# ---------------------------------------------------------------- ชนกับรายการที่เพิ่มเอง


@dataclass(frozen=True)
class EntrySnapshot:
    """ผู้เข้าสอบ 1 คนในรายชื่อที่ใช้อยู่ — เท่าที่ตัวสลับรายชื่อต้องรู้"""

    id: str
    candidate_no: str
    name_en: str | None
    name_th: str | None
    exam_mode: str
    school: str | None
    level: str | None
    source: str
    raw_award: str | None = None
    student_id: str | None = None

    @property
    def names(self) -> set[str]:
        return {n for n in (normalize_name(self.name_en or ""), normalize_name(self.name_th or "")) if n}


def row_names(row: RosterRow) -> set[str]:
    return {n for n in (normalize_name(row.name_en), normalize_name(row.name_th)) if n}


def detect_conflicts(rows: list[RosterRow], manual: list[EntrySnapshot]) -> list[dict[str, Any]]:
    """แถวใหม่ที่ชนกับรายการที่แอดมินเพิ่มเอง

    ชน = เลขเดียวกัน หรือชื่อเดียวกันแต่คนละเลข (น่าจะเป็นคนเดียวกันที่แอดมินใส่เลขผิดไว้)
    ระบบไม่ตัดสินเองว่าจะรวมหรือแยก เพราะทั้งสองทางเดาผิดได้ — ให้แอดมินเลือก
    """
    conflicts: list[dict[str, Any]] = []
    for entry in manual:
        for row in rows:
            if row.candidate_no == entry.candidate_no:
                reason = "SAME_NUMBER"
            elif row_names(row) & entry.names:
                reason = "SAME_NAME"
            else:
                continue
            conflicts.append({
                "id": f"{entry.id}:{row.candidate_no}",
                "reason": reason,
                "manualEntryId": entry.id,
                "incomingCandidateNo": row.candidate_no,
                "manual": {
                    "candidateNo": entry.candidate_no, "nameEn": entry.name_en, "nameTh": entry.name_th,
                    "examMode": entry.exam_mode, "school": entry.school, "level": entry.level,
                },
                "incoming": {
                    "row": row.row_number, "candidateNo": row.candidate_no,
                    "nameEn": row.name_en or None, "nameTh": row.name_th or None,
                    "examMode": row.exam_mode, "school": row.school or None, "level": row.level or None,
                },
            })
    return conflicts


# ---------------------------------------------------------------- วางแผนการสลับรายชื่อ

MERGE_FIELDS = ("candidateNo", "name", "examMode", "school", "level")


class ActivationRejected(ValueError):
    """สลับรายชื่อไม่ได้ — ข้อความบอกแอดมินได้ตรง ๆ ว่าต้องแก้อะไร"""


@dataclass
class ActivationPlan:
    creates: list[RosterRow] = field(default_factory=list)
    # (รายการ EXCEL เดิม, แถวใหม่ที่มาแทน)
    updates: list[tuple[EntrySnapshot, RosterRow]] = field(default_factory=list)
    deletes: list[EntrySnapshot] = field(default_factory=list)
    # (รายการที่เพิ่มเอง, ค่าหลังรวม, แถวที่รวมเข้ามา)
    merges: list[tuple[EntrySnapshot, dict[str, Any], RosterRow]] = field(default_factory=list)
    excluded_rows: list[RosterRow] = field(default_factory=list)


def plan_activation(
    existing: list[EntrySnapshot],
    rows: list[RosterRow],
    conflicts: list[dict[str, Any]],
    resolutions: list[dict[str, Any]],
) -> ActivationPlan:
    """วางแผนสลับรายชื่อ — ฟังก์ชันล้วน ทดสอบได้โดยไม่ต้องมีฐานข้อมูล

    กติกา:
      - รายการ EXCEL เดิมถูกแทนที่ทั้งหมดด้วยไฟล์ใหม่ (เลขเดิม = แก้ในแถวเดิม ไม่สร้างใหม่
        เพื่อให้ร่องรอยและการตัดสินที่ผูกกับคนนั้นอยู่ต่อได้)
      - รายการ MANUAL อยู่ต่อเสมอ ยกเว้นแอดมินเลือกรวมกับแถวใหม่
      - ทุกรายการที่ชนต้องมีคำตัดสิน และหลังตัดสินแล้วเลขผู้เข้าสอบต้องไม่ซ้ำกันเลย
    """
    by_id = {e.id: e for e in existing}
    rows_by_no = {r.candidate_no: r for r in rows}
    decided = {r.get("conflictId"): r for r in resolutions}

    plan = ActivationPlan()
    excluded: set[str] = set()
    merged: dict[str, str] = {}
    merge_values: dict[str, dict[str, Any]] = {}

    for conflict in conflicts:
        decision = decided.get(conflict["id"])
        if not decision:
            raise ActivationRejected("ยังมีรายการที่ชนกับผู้เข้าสอบที่เพิ่มเองซึ่งยังไม่ได้ตัดสิน")
        manual = by_id[conflict["manualEntryId"]]
        row = rows_by_no[conflict["incomingCandidateNo"]]
        if decision.get("action") == "KEEP_MANUAL":
            excluded.add(row.candidate_no)
        elif decision.get("action") == "MERGE":
            if manual.id in merge_values:
                raise ActivationRejected(
                    f"ผู้เข้าสอบที่เพิ่มเอง (เลข {manual.candidate_no}) รวมกับแถวใน Excel ได้แถวเดียว"
                )
            if row.candidate_no in merged:
                raise ActivationRejected(f"แถวเลข {row.candidate_no} ถูกเลือกให้รวมกับผู้เข้าสอบมากกว่าหนึ่งคน")
            merged[row.candidate_no] = manual.id
            merge_values[manual.id] = _merge(manual, row, decision.get("fields") or {})
            plan.merges.append((manual, merge_values[manual.id], row))
        else:
            raise ActivationRejected("คำตัดสินต้องเป็นรวม (MERGE) หรือเก็บรายการเดิม (KEEP_MANUAL)")

    both = excluded & merged.keys()
    if both:
        raise ActivationRejected(
            f"แถวเลข {', '.join(sorted(both))} ถูกตัดสินขัดกัน (ทั้งรวมและตัดทิ้ง)"
        )

    remaining = [r for r in rows if r.candidate_no not in excluded and r.candidate_no not in merged]
    plan.excluded_rows = [rows_by_no[no] for no in sorted(excluded)]

    # ตรวจเลขซ้ำของผลลัพธ์ทั้งชุด ก่อนแตะอะไรเลย
    final_numbers = [r.candidate_no for r in remaining]
    for entry in existing:
        if entry.source == "MANUAL":
            final_numbers.append(merge_values.get(entry.id, {}).get("candidateNo", entry.candidate_no))
    duplicates = sorted({n for n in final_numbers if final_numbers.count(n) > 1})
    if duplicates:
        raise ActivationRejected(
            f"หลังตัดสินแล้วเลขผู้เข้าสอบยังซ้ำกัน: {', '.join(duplicates)} — กรุณาเลือกใหม่"
        )

    excel = {e.candidate_no: e for e in existing if e.source == "EXCEL"}
    remaining_numbers = {r.candidate_no for r in remaining}
    plan.deletes = [e for no, e in excel.items() if no not in remaining_numbers]
    for row in remaining:
        if row.candidate_no in excel:
            plan.updates.append((excel[row.candidate_no], row))
        else:
            plan.creates.append(row)
    return plan


def _merge(manual: EntrySnapshot, row: RosterRow, fields: dict[str, str]) -> dict[str, Any]:
    """ค่าหลังรวม — แต่ละช่องแอดมินเลือกเองว่าใช้ของเดิม (MANUAL) หรือของใหม่ (INCOMING)
    ช่องที่ไม่ได้เลือกใช้ของใหม่ เพราะไฟล์ใหม่คือสิ่งที่ต้นทางยืนยันล่าสุด"""
    def pick(name: str, mine: Any, theirs: Any) -> Any:
        return mine if fields.get(name) == "MANUAL" else theirs

    use_manual_name = fields.get("name") == "MANUAL"
    return {
        "candidateNo": pick("candidateNo", manual.candidate_no, row.candidate_no),
        "nameEn": manual.name_en if use_manual_name else (row.name_en or None),
        "nameTh": manual.name_th if use_manual_name else (row.name_th or None),
        "examMode": pick("examMode", manual.exam_mode, row.exam_mode),
        "school": pick("school", manual.school, row.school or None),
        "level": pick("level", manual.level, row.level or None),
        "rawAward": row.award or manual.raw_award,
        "rowNumber": row.row_number,
    }


# ---------------------------------------------------------------- งานของ worker


def run_roster_validate(batch_id: str, on_progress: ProgressFn, payload: dict[str, Any]) -> dict[str, Any]:
    """ตรวจไฟล์รายชื่อที่เพิ่งอัป แล้วเก็บเป็นร่าง — ไม่แตะรายชื่อที่ใช้อยู่"""
    import_id = payload.get("importId")
    with connection() as conn:
        record = conn.execute(
            "SELECT id, source_key, status::text AS status FROM roster_imports WHERE id = %s AND batch_id = %s",
            (import_id, batch_id),
        ).fetchone()
    if record is None:
        raise ValueError("ไม่พบไฟล์รายชื่อที่ต้องตรวจ")
    if record["status"] not in ("PENDING", "READY"):
        # ร่างถูกทิ้งหรือถูกใช้ไปแล้วระหว่างรอคิว — ไม่ต้องทำอะไร
        return {"skipped": True, "status": record["status"]}

    validation = validate_roster(download_bytes(record["source_key"]))
    counts = validation.counts()
    conflicts = detect_conflicts(validation.rows, load_entries(batch_id, "MANUAL")) if validation.ok else []

    with connection() as conn, conn.transaction():
        conn.execute("DELETE FROM roster_import_rows WHERE import_id = %s", (import_id,))
        if validation.ok:
            with conn.cursor() as cur:
                cur.executemany(
                    """
                    INSERT INTO roster_import_rows
                      (id, import_id, row_number, candidate_no, name_en, name_th, exam_mode,
                       school, level, raw_award)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    [
                        (new_id(), import_id, r.row_number, r.candidate_no, r.name_en or None,
                         r.name_th or None, r.exam_mode, r.school or None, r.level or None,
                         r.award or None)
                        for r in validation.rows
                    ],
                )
        conn.execute(
            """
            UPDATE roster_imports
            SET status = %s, validated_at = NOW(), errors = %s, conflicts = %s,
                total_count = %s, online_count = %s, onsite_count = %s
            WHERE id = %s
            """,
            (
                "READY" if validation.ok else "INVALID",
                json.dumps(validation.errors, ensure_ascii=False),
                json.dumps(conflicts, ensure_ascii=False),
                counts["total"], counts["online"], counts["onsite"],
                import_id,
            ),
        )

    log.info("ตรวจรายชื่อ %s: %s แถว ปัญหา %s ชนกับที่เพิ่มเอง %s",
             import_id, counts["total"], len(validation.errors), len(conflicts))
    return {**counts, "errors": len(validation.errors), "conflicts": len(conflicts)}


def load_entries(batch_id: str, source: str | None = None, conn: Any = None) -> list[EntrySnapshot]:
    sql = """
        SELECT id::text, candidate_no, name_en, name_th, exam_mode::text AS exam_mode, school, level,
               source::text AS source, raw_award, student_id::text
        FROM roster_entries WHERE batch_id = %s
    """
    params: tuple = (batch_id,)
    if source:
        sql += " AND source = %s"
        params = (batch_id, source)
    # ลำดับต้องคงที่ รายการที่ชนจะได้เรียงเหมือนเดิมทุกครั้งที่ตรวจ
    sql += " ORDER BY candidate_no, id"

    def fetch(c: Any) -> list[EntrySnapshot]:
        return [EntrySnapshot(**row) for row in c.execute(sql, params).fetchall()]

    if conn is not None:
        return fetch(conn)
    with connection() as own:
        return fetch(own)


def run_roster_activate(batch_id: str, on_progress: ProgressFn, payload: dict[str, Any]) -> dict[str, Any]:
    """สลับรายชื่อที่ใช้อยู่เป็นชุดใหม่ในทรานแซกชันเดียว

    ตรวจรายการที่ชนกับที่เพิ่มเองใหม่อีกรอบ เพราะระหว่างที่รอแอดมินตัดสิน
    อาจมีคนเพิ่มหรือแก้รายการเองไปแล้ว — ถ้าไม่ตรงกับที่แอดมินเห็น ไม่สลับ ให้ตรวจใหม่
    """
    import_id = payload.get("importId")
    resolutions = payload.get("resolutions") or []
    session_id = payload.get("sessionId") or "system"

    try:
        with connection() as conn:
            with conn.transaction():
                stale, plan = _activate(conn, batch_id, import_id, resolutions, session_id)
    except Exception:
        # ไม่สำเร็จด้วยเหตุใดก็ตาม ร่างต้องกลับไปรอให้แอดมินกดใหม่ได้ ไม่ค้างอยู่ที่ "กำลังใช้"
        with connection() as conn:
            conn.execute(
                "UPDATE roster_imports SET status = 'READY' WHERE id = %s AND status = 'ACTIVATING'",
                (import_id,),
            )
        raise

    if stale:
        raise ValueError(
            "รายการที่ชนกับผู้เข้าสอบที่เพิ่มเองเปลี่ยนไปจากตอนที่ตรวจ (อาจมีการเพิ่มหรือแก้ไขระหว่างนั้น) "
            "กรุณาตรวจและตัดสินใหม่อีกครั้ง"
        )
    stats = {
        "created": len(plan.creates), "updated": len(plan.updates), "deleted": len(plan.deletes),
        "merged": len(plan.merges), "excluded": len(plan.excluded_rows),
    }
    log.info("สลับรายชื่อของ batch %s เป็น %s: %s", batch_id, import_id, stats)
    return stats


def _activate(
    conn: Any, batch_id: str, import_id: str, resolutions: list[dict[str, Any]], session_id: str
) -> tuple[bool, ActivationPlan]:
    batch = conn.execute(
        """
        SELECT b.id, b.status::text AS status, b.active_roster_import_id::text AS active_id,
               p.code AS program_code, e.round::text AS round, e.year
        FROM batches b JOIN exams e ON e.id = b.exam_id JOIN exam_programs p ON p.id = e.program_id
        WHERE b.id = %s FOR UPDATE OF b
        """,
        (batch_id,),
    ).fetchone()
    record = conn.execute(
        """
        SELECT id::text, status::text AS status, source_key, conflicts,
               total_count, online_count, onsite_count
        FROM roster_imports WHERE id = %s AND batch_id = %s FOR UPDATE
        """,
        (import_id, batch_id),
    ).fetchone()
    if record is None or record["status"] not in ("READY", "ACTIVATING"):
        raise ValueError("รายชื่อชุดนี้ใช้ไม่ได้แล้ว (อาจถูกทิ้งหรือมีชุดใหม่กว่า) กรุณาโหลดหน้าใหม่")
    if batch["status"] == "PUBLISHED":
        raise ValueError("รอบนี้เผยแพร่อยู่ ต้องยกเลิกการเผยแพร่ก่อนจึงจะเปลี่ยนรายชื่อได้")

    rows = _load_rows(conn, import_id)
    existing = load_entries(batch_id, conn=conn)
    conflicts = detect_conflicts(rows, [e for e in existing if e.source == "MANUAL"])
    if conflicts != (record["conflicts"] or []):
        conn.execute(
            "UPDATE roster_imports SET status = 'READY', conflicts = %s WHERE id = %s",
            (json.dumps(conflicts, ensure_ascii=False), import_id),
        )
        return True, ActivationPlan()

    plan = plan_activation(existing, rows, conflicts, resolutions)
    _apply_plan(conn, batch_id, import_id, plan)

    previous = batch["active_id"]
    if previous and previous != import_id:
        conn.execute("UPDATE roster_imports SET status = 'SUPERSEDED' WHERE id = %s", (previous,))
    conn.execute(
        "UPDATE roster_imports SET status = 'ACTIVE', activated_at = NOW(), resolutions = %s WHERE id = %s",
        (json.dumps(resolutions, ensure_ascii=False), import_id),
    )
    conn.execute(
        """
        UPDATE batches
        SET active_roster_import_id = %s, source_excel_key = %s,
            status = CASE WHEN status = 'DRAFT' THEN 'READY'::"BatchStatus" ELSE status END,
            updated_at = NOW()
        WHERE id = %s
        """,
        (import_id, record["source_key"], batch_id),
    )

    label = f"{batch['program_code']} {batch['round']} {batch['year']}"
    record_audit(
        conn, batch_id=batch_id, batch_label=label, entity_type="ROSTER_IMPORT", entity_id=import_id,
        action="ROSTER_ACTIVATED", session_id=session_id,
        before={"activeRosterImportId": previous},
        after={
            "activeRosterImportId": import_id,
            "total": record["total_count"], "online": record["online_count"],
            "onsite": record["onsite_count"], "created": len(plan.creates),
            "updated": len(plan.updates), "deleted": len(plan.deletes), "merged": len(plan.merges),
            "excludedRows": [r.candidate_no for r in plan.excluded_rows],
        },
    )
    decisions = {r.get("conflictId"): r for r in resolutions}
    for conflict in conflicts:
        decision = decisions[conflict["id"]]
        record_audit(
            conn, batch_id=batch_id, batch_label=label, entity_type="ROSTER_ENTRY",
            entity_id=conflict["manualEntryId"], action="ROSTER_CONFLICT_RESOLVED",
            session_id=session_id,
            before={"manual": conflict["manual"], "incoming": conflict["incoming"]},
            after={"action": decision.get("action"), "fields": decision.get("fields") or {}},
        )
    return False, plan


def _load_rows(conn: Any, import_id: str) -> list[RosterRow]:
    rows = conn.execute(
        """
        SELECT row_number, candidate_no, name_en, name_th, exam_mode::text AS exam_mode,
               level, school, raw_award
        FROM roster_import_rows WHERE import_id = %s ORDER BY row_number
        """,
        (import_id,),
    ).fetchall()
    return [
        RosterRow(
            row_number=r["row_number"], candidate_no=r["candidate_no"], name_en=r["name_en"] or "",
            name_th=r["name_th"] or "", exam_mode=r["exam_mode"], level=r["level"] or "",
            school=r["school"] or "", award=r["raw_award"] or "",
        )
        for r in rows
    ]


def _identity_fields(name_en: str | None, name_th: str | None, school: str | None) -> dict[str, Any]:
    return {
        "name_en_normalized": normalize_name(name_en or "") or None,
        "name_th_normalized": normalize_name(name_th or "") or None,
        "name_en_sort_key": name_sort_key(name_en or "") or None,
        "school_normalized": normalize_school(school or "") or None,
    }


def _apply_plan(conn: Any, batch_id: str, import_id: str, plan: ActivationPlan) -> None:
    """ลงมือตามแผน — ลำดับสำคัญเพราะเลขผู้เข้าสอบต้องไม่ซ้ำกันทุกขณะ:
    ลบรายการเดิมที่หายไปก่อน -> แก้รายการเดิม -> รวมรายการที่เพิ่มเอง -> สร้างรายการใหม่"""
    for entry in plan.deletes:
        # หน้าที่เคยผูกกับคนนี้กลับไปรอจับคู่ใหม่ (ใบที่ออกไปแล้วถูกลบตาม FK cascade)
        conn.execute(
            """
            UPDATE staging_pages
            SET roster_entry_id = NULL, manual_match = NULL, matched_manually = false,
                match_status = CASE WHEN match_status IN ('DISCARDED', 'SUPERSEDED', 'SKIPPED_FOREIGN',
                                                          'NATIONALITY_UNVERIFIED', 'PARSE_REVIEW')
                                    THEN match_status ELSE 'UNMATCHED' END
            WHERE roster_entry_id = %s
            """,
            (entry.id,),
        )
        conn.execute("DELETE FROM roster_entries WHERE id = %s", (entry.id,))

    for entry, row in plan.updates:
        identity = _identity_fields(row.name_en, row.name_th, row.school)
        names_changed = row_names(row) != entry.names
        conn.execute(
            """
            UPDATE roster_entries
            SET name_en = %s, name_th = %s, name_en_normalized = %s, name_th_normalized = %s,
                name_en_sort_key = %s, exam_mode = %s, school = %s, school_normalized = %s,
                level = %s, raw_award = %s, source_row = %s, roster_import_id = %s,
                student_id = CASE WHEN %s THEN NULL ELSE student_id END,
                student_linked_manually = CASE WHEN %s THEN false ELSE student_linked_manually END,
                version = version + 1, updated_at = NOW()
            WHERE id = %s
            """,
            (
                row.name_en or None, row.name_th or None, identity["name_en_normalized"],
                identity["name_th_normalized"], identity["name_en_sort_key"], row.exam_mode,
                row.school or None, identity["school_normalized"], row.level or None,
                row.award or None, row.row_number, import_id,
                # ชื่อเปลี่ยน = ตัวคนที่ผูกไว้อาจไม่ใช่คนนี้แล้ว ต้องระบุตัวใหม่
                names_changed, names_changed, entry.id,
            ),
        )

    for manual, values, _row in plan.merges:
        identity = _identity_fields(values["nameEn"], values["nameTh"], values["school"])
        names_changed = {n for n in (identity["name_en_normalized"], identity["name_th_normalized"]) if n} != manual.names
        conn.execute(
            """
            UPDATE roster_entries
            SET candidate_no = %s, name_en = %s, name_th = %s, name_en_normalized = %s,
                name_th_normalized = %s, name_en_sort_key = %s, exam_mode = %s, school = %s,
                school_normalized = %s, level = %s, raw_award = %s, source_row = %s,
                source = 'EXCEL', roster_import_id = %s,
                student_id = CASE WHEN %s THEN NULL ELSE student_id END,
                student_linked_manually = CASE WHEN %s THEN false ELSE student_linked_manually END,
                version = version + 1, updated_at = NOW()
            WHERE id = %s
            """,
            (
                values["candidateNo"], values["nameEn"], values["nameTh"],
                identity["name_en_normalized"], identity["name_th_normalized"],
                identity["name_en_sort_key"], values["examMode"], values["school"],
                identity["school_normalized"], values["level"], values["rawAward"],
                values["rowNumber"], import_id, names_changed, names_changed, manual.id,
            ),
        )

    if plan.creates:
        values = []
        for row in plan.creates:
            identity = _identity_fields(row.name_en, row.name_th, row.school)
            values.append((
                new_id(), batch_id, row.candidate_no, row.name_en or None, row.name_th or None,
                identity["name_en_normalized"], identity["name_th_normalized"],
                identity["name_en_sort_key"], row.exam_mode, row.school or None,
                identity["school_normalized"], row.level or None, row.award or None,
                row.row_number, import_id,
            ))
        with conn.cursor() as cur:
            cur.executemany(
                """
                INSERT INTO roster_entries
                  (id, batch_id, candidate_no, name_en, name_th, name_en_normalized,
                   name_th_normalized, name_en_sort_key, exam_mode, source, school,
                   school_normalized, level, raw_award, source_row, roster_import_id, updated_at)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, 'EXCEL', %s, %s, %s, %s, %s, %s, NOW())
                """,
                values,
            )
