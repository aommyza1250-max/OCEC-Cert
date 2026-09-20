"""จับคู่รายชื่อจาก Excel เข้ากับหน้าที่ตัดแยกไว้แล้ว แล้วบันทึกเป็นเกียรติบัตรจริง

**วนจาก "หน้าที่ตัดได้" ไม่ใช่ "แถวใน Excel"**
เพราะ Excel 1 แถวรองรับได้หลายหน้า — ของจริงคนที่ได้ Perfect Score จะมีทั้งหน้า Gold
และหน้า Perfect Score แต่มีแถวเดียวใน Excel เขียนว่า PERFECT SCORER
ถ้าวนจากแถว หน้า Gold ของคนเหล่านี้จะหลุดหายไปเลย

ลำดับความมั่นใจในการจับคู่:
  1. เลขบนหน้า (Cert No) ตรงกับ CANDIDATE NO ใน Excel  <- ของจริงตรงกัน 242/242
  2. ชื่อที่ normalize แล้วตรงกับแถวเดียว

**รางวัลของใบที่ออกมาจากชื่อโฟลเดอร์ใน ZIP เสมอ ไม่ใช่จากคอลัมน์ AWARD ใน Excel**
เพราะ Excel บันทึกรางวัลสูงสุดของคนนั้นแค่รางวัลเดียว

หลักการสำคัญ: **ถ้าไม่มั่นใจ ห้ามเดา** อะไรที่ระบุตัวไม่ได้ให้ส่งต่อให้แอดมินดูใบจริง
ดีกว่าจับคู่ผิดแล้วผู้ปกครองโหลดได้เกียรติบัตรของคนอื่น
"""

import io
import logging
import re
from dataclasses import dataclass
from typing import Any, Callable

from openpyxl import load_workbook

from ..db import connection, new_id
from ..normalize import name_sort_key, normalize_award, normalize_name, normalize_school
from ..storage import download_bytes

log = logging.getLogger(__name__)

ProgressFn = Callable[[dict[str, Any]], None]

# หัวคอลัมน์ที่ยอมรับ — เทียบหลังตัดช่องว่าง จุด ขีดล่าง และทำตัวพิมพ์เล็กแล้ว
HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    # ของจริงใช้ CANDIDATE NO ซึ่งเป็นเลขเดียวกับ "Cert No" ที่พิมพ์บนหน้าเกียรติบัตร
    "cert_no": ("candidateno", "candidatenumber", "certno", "certificateno", "certificatenumber",
                "เลขที่นั่งสอบ", "เลขประจำตัวสอบ", "รหัสผู้เข้าสอบ", "เลขเกียรติบัตร",
                "เลขที่เกียรติบัตร"),
    "name_en": ("candidatename", "name", "fullname", "englishname", "nameenglish", "nameen",
                "nameeng", "ชื่อภาษาอังกฤษ", "ชื่ออังกฤษ", "ชื่อ-นามสกุลภาษาอังกฤษ"),
    "first_en": ("firstname", "givenname", "given"),
    "last_en": ("lastname", "surname", "familyname", "family"),
    "name_th": ("ชื่อ-นามสกุล", "ชื่อ-สกุล", "ชื่อนามสกุล", "ชื่อภาษาไทย", "ชื่อไทย",
                "ชื่อ-นามสกุลภาษาไทย", "thainame"),
    "first_th": ("ชื่อ", "ชื่อจริง"),
    "last_th": ("นามสกุล", "สกุล"),
    "level": ("grade", "ระดับชั้น", "ชั้น", "class"),
    "award": ("รางวัล", "ผลการแข่งขัน", "ผลรางวัล", "award", "result", "medal", "prize"),
    "school": ("โรงเรียน", "สถานศึกษา", "ชื่อโรงเรียน", "school", "schoolname", "institution"),
}

MAX_HEADER_SCAN_ROWS = 10
# จำนวนรายการที่เก็บรายละเอียดไว้ใน stats — เกินกว่านี้เก็บแค่ตัวเลข
MAX_REPORTED = 200


@dataclass
class RosterRow:
    row_number: int
    cert_no: str
    name_en: str
    name_th: str
    level: str
    school: str
    award: str


def run_match(batch_id: str, on_progress: ProgressFn) -> dict[str, Any]:
    batch = _load_batch(batch_id)
    if not batch["source_excel_key"]:
        raise ValueError("batch นี้ยังไม่มีไฟล์ Excel รายชื่อ")

    rows = parse_roster(download_bytes(batch["source_excel_key"]))
    log.info("อ่านรายชื่อจาก Excel ได้ %s แถว", len(rows))

    by_cert_no = {r.cert_no: r for r in rows if r.cert_no}
    by_name: dict[str, list[RosterRow]] = {}
    for row in rows:
        for key in {normalize_name(row.name_en), normalize_name(row.name_th)}:
            if key:
                by_name.setdefault(key, []).append(row)

    # จำสถานะเผยแพร่เดิมไว้ก่อนล้าง ไม่งั้นรันจับคู่ซ้ำทีไร ผู้ปกครองจะค้นไม่เจอจนกว่าจะกดเผยแพร่ใหม่
    published = _published_snapshot(batch_id)

    # รันจับคู่ซ้ำได้: ล้างผลอัตโนมัติรอบก่อนทิ้งก่อนเสมอ
    # (แอดมินอัปโหลด Excel ใหม่ทับได้บ่อย) แต่เก็บสิ่งที่แอดมินตัดสินด้วยมือไว้
    _reset_previous_matches(batch_id)

    pages = _load_pages(batch_id)
    stats = _new_stats(len(rows), len(pages))

    # ภายในรอบนำเข้าเดียวกัน เลขเดียวกัน = คนเดียวกัน ไม่ต้องเดาจากชื่อ
    student_by_cert: dict[str, str] = {}
    # (ผู้เข้าสอบ, รางวัล) ที่ออกใบไปแล้วในรอบนี้ — ซ้ำคู่นี้แปลว่ามีอะไรผิด
    # เริ่มจากใบที่แอดมินจับคู่ด้วยมือไว้แล้ว เพราะเราไม่แตะของพวกนั้น
    issued: set[tuple[str, str]] = _manual_pairs(batch_id)
    matched_rows: set[int] = set()

    for index, page in enumerate(pages, start=1):
        row, how = _find_row(page, by_cert_no, by_name)

        if row is None:
            _mark(page["id"], "UNMATCHED", None)
            stats["pagesUnmatched"] += 1
        elif how == "cert" and not _names_agree(page, row):
            _mark(page["id"], "AMBIGUOUS",
                  f"เลข {page['cert_no']} ตรงกับ Excel แถวที่ {row.row_number} "
                  f"แต่ชื่อไม่ตรงกัน: บนเกียรติบัตรเขียน '{page['extracted_name']}' "
                  f"ส่วน Excel เขียน '{row.name_en or row.name_th}'",
                  roster_award=row.award)
            stats["nameMismatch"] += 1
        elif how == "name_many":
            _mark(page["id"], "AMBIGUOUS",
                  f"ชื่อนี้ตรงกับ Excel มากกว่าหนึ่งแถว และหน้านี้ไม่มีเลขให้ยืนยัน", None)
            stats["ambiguous"] += 1
        else:
            student_id = _resolve_student(page, row, student_by_cert)
            if student_id is None:
                _mark(page["id"], "AMBIGUOUS",
                      f"มีผู้เข้าสอบชื่อ '{row.name_en or row.name_th}' อยู่ในระบบมากกว่าหนึ่งคน "
                      "และข้อมูลที่มีแยกไม่ออกว่าเป็นคนไหน "
                      "กรุณาจับคู่ด้วยมือ โดยระบุโรงเรียนเพื่อให้รอบหน้าระบบแยกได้เอง",
                      roster_award=row.award)
                stats["ambiguous"] += 1
                continue
            key = (student_id, page["award"] or "")
            if key in issued:
                _mark(page["id"], "DUPLICATE_NAME",
                      f"ผู้เข้าสอบคนนี้ได้รางวัล {page['award']} ไปแล้วจากหน้าอื่นในรอบนี้ "
                      "กรุณาเทียบเกียรติบัตรทั้งสองใบว่าเป็นคนละคน หรือเป็นไฟล์ซ้ำ",
                      roster_award=row.award)
                stats["duplicateNames"] += 1
            else:
                issued.add(key)
                matched_rows.add(row.row_number)
                _commit_match(batch, page, row, student_id, published)
                stats["matched"] += 1
                stats["matchedByCertNo" if how == "cert" else "matchedByName"] += 1
                _cross_check(stats, page, row)

        if index % 25 == 0 or index == len(pages):
            on_progress({"stage": "match", "done": index, "total": len(pages)})

    unused = [r for r in rows if r.row_number not in matched_rows]
    stats["rowsNotUsed"] = len(unused)
    stats["unmatchedRows"] = [
        {
            "row": r.row_number,
            "certNo": r.cert_no,
            "name": r.name_en or r.name_th,
            # รางวัลที่ควรจะเป็น เอาไปบอกระบบตอนแอดมินอัปไฟล์ที่ขาดเข้ามาทีหลัง
            "award": normalize_award(r.award) or None,
        }
        for r in unused[:MAX_REPORTED]
    ]

    log.info("จับคู่ batch %s เสร็จ: %s", batch_id, {k: v for k, v in stats.items()
                                                    if k != "unmatchedRows"})
    return stats


def _new_stats(row_count: int, page_count: int) -> dict[str, Any]:
    return {
        "rosterRows": row_count,
        "pagesToMatch": page_count,
        "matched": 0,
        "matchedByCertNo": 0,
        "matchedByName": 0,
        "nameMismatch": 0,
        "ambiguous": 0,
        "duplicateNames": 0,
        "pagesUnmatched": 0,
        "rowsNotUsed": 0,
        "awardMismatchWithRoster": 0,
        "levelMismatch": 0,
        "unmatchedRows": [],
    }


# ---------------------------------------------------------------- หาแถวที่คู่กัน

def _find_row(
    page: dict[str, Any],
    by_cert_no: dict[str, RosterRow],
    by_name: dict[str, list[RosterRow]],
) -> tuple[RosterRow | None, str]:
    """คืน (แถวที่คู่กัน, วิธีที่หาเจอ) — วิธีเป็น 'cert' / 'name' / 'name_many' / ''"""
    cert_no = page.get("cert_no")
    if cert_no and cert_no in by_cert_no:
        return by_cert_no[cert_no], "cert"

    key = page.get("extracted_name_normalized")
    if key:
        candidates = by_name.get(key, [])
        if len(candidates) == 1:
            return candidates[0], "name"
        if len(candidates) > 1:
            return candidates[0], "name_many"
    return None, ""


def _names_agree(page: dict[str, Any], row: RosterRow) -> bool:
    """ชื่อบนหน้ากับชื่อใน Excel ต้องตรงกันเพื่อยืนยันว่าเลขไม่ได้ชนกันโดยบังเอิญ

    ถ้าหน้านั้นอ่านชื่อไม่ออกเลย ให้เชื่อเลขไปก่อน — ดีกว่าทิ้งใบนั้นไปเฉย ๆ
    """
    page_name = page.get("extracted_name_normalized")
    if not page_name:
        return True
    return page_name in {normalize_name(row.name_en), normalize_name(row.name_th)}


def _cross_check(stats: dict[str, Any], page: dict[str, Any], row: RosterRow) -> None:
    """เทียบข้อมูลที่ควรตรงกันแต่ไม่ถึงกับทำให้จับคู่ไม่ได้ — นับไว้ให้แอดมินเห็น"""
    roster_award = normalize_award(row.award)
    if roster_award and page.get("award") and roster_award != page["award"]:
        stats["awardMismatchWithRoster"] += 1
    if row.level and page.get("level") and row.level.strip().upper() != page["level"].strip().upper():
        stats["levelMismatch"] += 1


# ---------------------------------------------------------------- อ่าน Excel

def parse_roster(xlsx_bytes: bytes) -> list[RosterRow]:
    workbook = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    sheet = workbook.active

    grid = [list(r) for r in sheet.iter_rows(values_only=True)]
    header_index, columns = _find_header(grid)
    if header_index is None:
        raise ValueError(
            "หาหัวตารางในไฟล์ Excel ไม่เจอ — ต้องมีคอลัมน์ชื่ออย่างน้อย 1 คอลัมน์ "
            f"(ที่รองรับ: {', '.join(sorted(HEADER_ALIASES))})"
        )

    rows: list[RosterRow] = []
    for offset, raw in enumerate(grid[header_index + 1:], start=header_index + 2):
        row = _build_row(raw, columns, offset)
        # ข้ามแถวว่างและแถวรวมยอดท้ายตาราง
        if row and (row.name_en or row.name_th):
            rows.append(row)
    return rows


def _find_header(grid: list[list[Any]]) -> tuple[int | None, dict[str, int]]:
    for index, raw in enumerate(grid[:MAX_HEADER_SCAN_ROWS]):
        columns: dict[str, int] = {}
        for col, cell in enumerate(raw):
            field = _match_header(cell)
            # คอลัมน์แรกที่เจอชนะ กันกรณีมีหัวซ้ำ
            if field and field not in columns:
                columns[field] = col
        if _has_name_column(columns):
            return index, columns
    return None, {}


def _match_header(cell: Any) -> str | None:
    if cell is None:
        return None
    key = "".join(str(cell).split()).lower().replace("_", "").replace(".", "")
    if not key:
        return None
    for field, aliases in HEADER_ALIASES.items():
        if key in aliases:
            return field
    return None


def _has_name_column(columns: dict[str, int]) -> bool:
    return bool(
        {"name_en", "name_th"} & columns.keys()
        or {"first_en", "last_en"} <= columns.keys()
        or {"first_th", "last_th"} <= columns.keys()
    )


def _build_row(raw: list[Any], columns: dict[str, int], row_number: int) -> RosterRow | None:
    def value(field: str) -> str:
        col = columns.get(field)
        if col is None or col >= len(raw) or raw[col] is None:
            return ""
        return str(raw[col]).strip()

    name_en = value("name_en") or " ".join(x for x in (value("first_en"), value("last_en")) if x)
    name_th = value("name_th") or " ".join(x for x in (value("first_th"), value("last_th")) if x)

    if not name_en and not name_th:
        return None
    return RosterRow(
        row_number=row_number,
        cert_no=_digits_only(value("cert_no")),
        name_en=name_en,
        name_th=name_th,
        level=value("level"),
        school=value("school"),
        award=value("award"),
    )


def _digits_only(text: str) -> str:
    """ตัดให้เหลือแต่ตัวเลข

    ต้องตัดส่วนทศนิยมทิ้งก่อน เพราะ openpyxl อ่านเซลล์ตัวเลขมาเป็น float
    แล้ว str() ได้ '900101.0' — ถ้าเก็บแค่ตัวเลขตรง ๆ จะกลายเป็น '9001010'
    ซึ่งทำให้เลขผู้เข้าสอบเพี้ยนทั้งไฟล์โดยไม่มีอะไรฟ้อง
    """
    return "".join(ch for ch in re.sub(r"\.\d+$", "", text.strip()) if ch.isdigit())


# ---------------------------------------------------------------- เขียนฐานข้อมูล

def _resolve_student(
    page: dict[str, Any],
    row: RosterRow,
    student_by_cert: dict[str, str],
) -> str | None:
    """หาว่าหน้านี้เป็นของผู้เข้าสอบคนไหน — คืน None ถ้าระบุตัวไม่ได้

    ภายในรอบนำเข้าเดียวกัน เลขผู้เข้าสอบเดียวกัน = คนเดียวกันแน่นอน ไม่ต้องเดาจากชื่อ
    ข้ามรอบ/ข้ามปีไม่มีเลขให้อ้าง จึงยังต้องใช้ชื่อ (+โรงเรียนถ้ามี) รวมคนเดิมเข้าด้วยกัน
    ซึ่งเป็นสิ่งที่ทำให้หน้าค้นหารวมเกียรติบัตรทุกใบของคนคนนั้นไว้ที่เดียวได้

    ถ้ามีคนชื่อเดียวกันในระบบหลายคนและแยกไม่ออก **ห้ามสร้างคนใหม่**
    เพราะการสร้างคนใหม่ก็เป็นการเดาอย่างหนึ่ง (เดาว่า "เป็นคนละคน") และถ้ารันจับคู่ซ้ำ
    จะเกิดผู้เข้าสอบซ้ำซ้อนขึ้นเรื่อย ๆ โดยไม่มีอะไรฟ้อง — ส่งให้แอดมินตัดสินแทน
    """
    cert_no = row.cert_no or page.get("cert_no") or ""
    if cert_no and cert_no in student_by_cert:
        return student_by_cert[cert_no]

    name_en_norm = normalize_name(row.name_en) or None
    name_th_norm = normalize_name(row.name_th) or None
    school_norm = normalize_school(row.school) or None

    with connection() as conn:
        # ต้อง cast ::text ตรง ๆ ไม่งั้น Postgres เดาชนิดของพารามิเตอร์ใน 'IS NOT NULL' ไม่ออก
        candidates = conn.execute(
            """
            SELECT id, school_normalized FROM students
            WHERE (%s::text IS NOT NULL AND name_en_normalized = %s)
               OR (%s::text IS NOT NULL AND name_th_normalized = %s)
            ORDER BY created_at
            """,
            (name_en_norm, name_en_norm, name_th_norm, name_th_norm),
        ).fetchall()

        chosen = _narrow_by_school(candidates, school_norm)
        if len(chosen) > 1:
            return None
        if len(chosen) == 1:
            student_id = chosen[0]["id"]
            _fill_missing_fields(conn, student_id, row, name_th_norm, name_en_norm, school_norm)
        else:
            student_id = _create_student(conn, row, name_th_norm, name_en_norm, school_norm)

    if cert_no:
        student_by_cert[cert_no] = student_id
    return student_id


def _narrow_by_school(
    candidates: list[dict[str, Any]], school_norm: str | None
) -> list[dict[str, Any]]:
    """คัดผู้เข้าสอบชื่อพ้องให้เหลือเฉพาะคนที่อยู่โรงเรียนเดียวกัน

    - เจอคนโรงเรียนเดียวกัน -> เอาเฉพาะกลุ่มนั้น
    - ไม่เจอ แต่มีคนที่ยังไม่เคยบันทึกโรงเรียน -> ถือว่าน่าจะใช่ แล้วค่อยเติมโรงเรียนให้
    - ทุกคนที่ชื่อนี้อยู่คนละโรงเรียน -> เป็นคนใหม่แน่นอน
    """
    if not school_norm or not candidates:
        return candidates

    same_school = [c for c in candidates if c["school_normalized"] == school_norm]
    if same_school:
        return same_school
    return [c for c in candidates if not c["school_normalized"]]


def _commit_match(
    batch: dict[str, Any],
    page: dict[str, Any],
    row: RosterRow,
    student_id: str,
    published: dict[tuple[str, str], Any],
) -> None:
    with connection() as conn:
        with conn.transaction():
            conn.execute(
                """
                UPDATE staging_pages
                SET match_status = 'MATCHED', matched_student_id = %s,
                    matched_manually = false, match_note = NULL, roster_award = %s
                WHERE id = %s
                """,
                (student_id, row.award or None, page["id"]),
            )
            conn.execute(
                """
                INSERT INTO certificates
                  (id, student_id, exam_id, batch_id, staging_page_id,
                   pdf_key, preview_key, page_number, award, cert_no, candidate_no, level,
                   published_at)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON CONFLICT (exam_id, student_id, award) DO UPDATE
                SET staging_page_id = EXCLUDED.staging_page_id,
                    batch_id = EXCLUDED.batch_id,
                    pdf_key = EXCLUDED.pdf_key,
                    preview_key = EXCLUDED.preview_key,
                    page_number = EXCLUDED.page_number,
                    cert_no = EXCLUDED.cert_no,
                    candidate_no = EXCLUDED.candidate_no,
                    level = EXCLUDED.level
                """,
                (
                    new_id(),
                    student_id,
                    batch["exam_id"],
                    batch["id"],
                    page["id"],
                    page["pdf_key"],
                    page["preview_key"],
                    page["page_number"],
                    # รางวัลมาจากโฟลเดอร์ใน ZIP ไม่ใช่จาก Excel
                    page["award"],
                    page["cert_no"],
                    row.cert_no or None,
                    # ระดับชั้นอ่านจากหน้ากระดาษก่อน เพราะนั่นคือสิ่งที่ผู้ปกครองถืออยู่ในมือ
                    page["level"] or row.level or None,
                    # คงสถานะเผยแพร่เดิมไว้ ไม่ตัดสินใหม่เอง
                    # การตัดสินว่าใบไหนเผยแพร่มีที่เดียวคือปุ่มเผยแพร่ฝั่งเว็บ (apps/web/src/lib/publish.ts)
                    # ถ้าตัวจับคู่มาตั้งเองด้วย กฎจะเหลื่อมกันแล้วใบที่ควรถูกกันไว้จะหลุดออกไปเงียบ ๆ
                    published.get((page["cert_no"], page["award"]))
                    or published.get((student_id, page["award"])),
                ),
            )


def _fill_missing_fields(
    conn: Any,
    student_id: str,
    row: RosterRow,
    name_th_norm: str | None,
    name_en_norm: str | None,
    school_norm: str | None,
) -> None:
    """เติมข้อมูลที่เดิมยังว่างอยู่ (Excel คนละรอบอาจมีคนละภาษา หรือเพิ่งเริ่มมีคอลัมน์โรงเรียน)"""
    conn.execute(
        """
        UPDATE students
        SET name_th = COALESCE(name_th, %s),
            name_en = COALESCE(name_en, %s),
            name_th_normalized = COALESCE(name_th_normalized, %s),
            name_en_normalized = COALESCE(name_en_normalized, %s),
            name_en_sort_key = COALESCE(name_en_sort_key, %s),
            school = COALESCE(school, %s),
            school_normalized = COALESCE(school_normalized, %s)
        WHERE id = %s
        """,
        (
            row.name_th or None,
            row.name_en or None,
            name_th_norm,
            name_en_norm,
            name_sort_key(row.name_en) or None,
            row.school or None,
            school_norm,
            student_id,
        ),
    )


def _create_student(
    conn: Any,
    row: RosterRow,
    name_th_norm: str | None,
    name_en_norm: str | None,
    school_norm: str | None,
) -> str:
    student_id = new_id()
    conn.execute(
        """
        INSERT INTO students
          (id, name_th, name_en, name_th_normalized, name_en_normalized, name_en_sort_key,
           school, school_normalized)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        """,
        (
            student_id,
            row.name_th or None,
            row.name_en or None,
            name_th_norm,
            name_en_norm,
            name_sort_key(row.name_en) or None,
            row.school or None,
            school_norm,
        ),
    )
    return student_id


def _mark(page_id: str, status: str, note: str | None, roster_award: str | None = None) -> None:
    with connection() as conn:
        conn.execute(
            """
            UPDATE staging_pages
            SET match_status = %s, match_note = %s, roster_award = COALESCE(%s, roster_award)
            WHERE id = %s AND match_status <> 'MATCHED'
            """,
            (status, note, roster_award, page_id),
        )


def _published_snapshot(batch_id: str) -> dict[tuple[str, str], Any]:
    """สถานะเผยแพร่ของเกียรติบัตรที่มีอยู่ ก่อนจะถูกลบแล้วสร้างใหม่

    ทำคีย์ไว้สองแบบ: (เลขผู้เข้าสอบ, รางวัล) ซึ่งเสถียรที่สุด
    และ (ผู้เข้าสอบ, รางวัล) เผื่อหน้าที่อ่านเลขไม่ได้
    """
    with connection() as conn:
        rows = conn.execute(
            """
            SELECT cert_no, student_id, award, published_at
            FROM certificates WHERE batch_id = %s AND published_at IS NOT NULL
            """,
            (batch_id,),
        ).fetchall()

    snapshot: dict[tuple[str, str], Any] = {}
    for row in rows:
        if row["cert_no"]:
            snapshot[(row["cert_no"], row["award"])] = row["published_at"]
        snapshot[(row["student_id"], row["award"])] = row["published_at"]
    return snapshot


def _reset_previous_matches(batch_id: str) -> None:
    """ล้างผลจับคู่อัตโนมัติของรอบนำเข้านี้ เพื่อให้รันใหม่ได้ผลเหมือนเริ่มต้นใหม่

    ไม่แตะหน้าที่ `matched_manually = true` เพราะนั่นคือการตัดสินของคน
    ซึ่งไม่ควรถูกลบทิ้งเพราะแอดมินอัปโหลด Excel ใหม่
    """
    with connection() as conn:
        conn.execute(
            """
            DELETE FROM certificates c
            USING staging_pages sp
            WHERE c.staging_page_id = sp.id
              AND c.batch_id = %s
              AND sp.matched_manually = false
            """,
            (batch_id,),
        )
        conn.execute(
            """
            UPDATE staging_pages
            SET match_status = 'UNMATCHED', matched_student_id = NULL, match_note = NULL
            WHERE batch_id = %s
              AND matched_manually = false
              AND match_status NOT IN ('SKIPPED_FOREIGN', 'DISCARDED')
            """,
            (batch_id,),
        )


def _manual_pairs(batch_id: str) -> set[tuple[str, str]]:
    """คู่ (ผู้เข้าสอบ, รางวัล) ที่แอดมินจับไว้ด้วยมือแล้วในรอบนำเข้านี้"""
    with connection() as conn:
        rows = conn.execute(
            "SELECT student_id, award FROM certificates WHERE batch_id = %s", (batch_id,)
        ).fetchall()
    return {(r["student_id"], r["award"]) for r in rows}


def _load_batch(batch_id: str) -> dict[str, Any]:
    with connection() as conn:
        row = conn.execute(
            """
            SELECT b.id, b.exam_id, b.source_excel_key,
                   p.code AS program_code, e.round, e.year
            FROM batches b
            JOIN exams e ON e.id = b.exam_id
            JOIN exam_programs p ON p.id = e.program_id
            WHERE b.id = %s
            """,
            (batch_id,),
        ).fetchone()
    if row is None:
        raise ValueError(f"ไม่พบ batch {batch_id}")
    return row


def _load_pages(batch_id: str) -> list[dict[str, Any]]:
    with connection() as conn:
        return conn.execute(
            """
            SELECT id, page_number, pdf_key, preview_key, cert_no, level, award,
                   extracted_name, extracted_name_normalized
            FROM staging_pages
            WHERE batch_id = %s
              AND matched_manually = false
              AND match_status NOT IN ('SKIPPED_FOREIGN', 'DISCARDED')
            ORDER BY page_number
            """,
            (batch_id,),
        ).fetchall()
